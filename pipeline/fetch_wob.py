"""EC Weekly Oil Bulletin: weekly retail prices (with and without taxes), VAT history, annual consumption."""
import datetime as dt

import openpyxl

from common import CACHE, EU27, FUELS, WOB_URL, http_get

AREAS = EU27 + ["EU"]


def download(force: bool = True):
    CACHE.mkdir(parents=True, exist_ok=True)
    path = CACHE / "wob_history.xlsx"
    if force or not path.exists():
        path.write_bytes(http_get(WOB_URL, timeout=300))
    return path


def _prices(ws, kind: str) -> dict:
    """{area: {fuel: {date: EUR per litre}}} from a 'Prices with/wo taxes' sheet (values are EUR/1000 L)."""
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    col = {h: i for i, h in enumerate(header) if isinstance(h, str)}
    out = {a: {f: {} for f in FUELS} for a in AREAS}
    for r in rows:
        d = r[0]
        if not isinstance(d, dt.datetime):
            continue
        for a in AREAS:
            for f, (suffix, _, _) in FUELS.items():
                v = r[col[f"{a}_price_{kind}_{suffix}"]]
                if isinstance(v, (int, float)) and v > 0:
                    out[a][f][d.date()] = v / 1000.0
    return out


def _vat(ws) -> dict:
    """{area: [(since_date, {fuel: vat_pct})]} sorted newest first."""
    out, cur = {}, None
    for r in list(ws.iter_rows(values_only=True))[4:]:
        if isinstance(r[0], str) and r[0].endswith("_"):
            cur = r[0][:-1]
            out.setdefault(cur, [])
        if cur is None or not isinstance(r[1], dt.datetime):
            continue
        out[cur].append((r[1].date(), {"petrol": r[2], "diesel": r[3]}))
    for k in out:
        out[k].sort(key=lambda x: x[0], reverse=True)
    return out


def _consumption(ws) -> dict:
    """Annual consumption in kt: {year: {area: {fuel: kt}}}. Country columns only (EU cells are formulas)."""
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    col = {h: i for i, h in enumerate(header) if isinstance(h, str)}
    out = {}
    for r in rows[2:]:
        if not isinstance(r[0], int):
            continue
        out[r[0]] = {}
        for a in EU27:
            out[r[0]][a] = {}
            for f, (suffix, _, _) in FUELS.items():
                v = r[col.get(f"{a}_consumption_{suffix}", -1)] if f"{a}_consumption_{suffix}" in col else None
                if isinstance(v, (int, float)):
                    out[r[0]][a][f] = float(v)
    return out


def vat_at(vat_hist: dict, area: str, fuel: str, day: dt.date) -> float:
    for since, rates in vat_hist.get(area, []):
        if since <= day and rates.get(fuel) is not None:
            return float(rates[fuel]) / 100.0
    return 0.20


def load(force_download: bool = True) -> dict:
    path = download(force_download)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data = {
        "with_tax": _prices(wb["Prices with taxes"], "with_tax"),
        "wo_tax": _prices(wb["Prices wo taxes"], "wo_tax"),
        "vat": _vat(wb["VAT"]),
        "consumption": _consumption(wb["Consumption"]),
    }
    wb.close()
    return data


if __name__ == "__main__":
    d = load(force_download=False)
    eu = d["with_tax"]["EU"]["petrol"]
    last = max(eu)
    print("latest bulletin", last, "EU petrol", round(eu[last], 3), "EU diesel", round(d["with_tax"]["EU"]["diesel"][last], 3))
    print("VAT ES diesel on 2026-04-01:", vat_at(d["vat"], "ES", "diesel", dt.date(2026, 4, 1)))
    print("consumption years", sorted(d["consumption"])[-3:], "IT", d["consumption"][max(d["consumption"])]["IT"])
